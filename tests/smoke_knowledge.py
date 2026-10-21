"""Opt-in localhost integration/resource probe using synthetic documents only."""
import asyncio
import ctypes
from datetime import datetime
from io import BytesIO
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import threading
import time

ROOT = Path('D:/ITGlasses')
sys.path.insert(0, str(ROOT))
os.environ['TEMP'] = os.environ['TMP'] = str(ROOT / '.cache/tmp')
import httpx
from backend.app.knowledge.embeddings import OllamaEmbeddings
from backend.app.knowledge.grounding import EvidenceSelector
from backend.app.knowledge.service import KnowledgeService

if len(sys.argv) == 3 and sys.argv[1] == '--serve':
    import uvicorn
    from backend.app.main import create_app
    svc = KnowledgeService(Path(sys.argv[2]), OllamaEmbeddings(), EvidenceSelector('qwen3:4b-instruct-2507-q4_K_M'))
    uvicorn.run(create_app(knowledge=svc), host='127.0.0.1', port=8001, access_log=False, log_level='critical')
    raise SystemExit()

class Memory(ctypes.Structure):
    _fields_ = [('length', ctypes.c_ulong), ('load', ctypes.c_ulong)] + [(name, ctypes.c_ulonglong) for name in ['total','available','page_total','page_available','virtual_total','virtual_available','extended']]

phase = 'baseline'
samples = []
stop = threading.Event()

def measure():
    memory = Memory(); memory.length = ctypes.sizeof(memory)
    ctypes.windll.kernel32.GlobalMemoryStatusEx(ctypes.byref(memory))
    gpu = subprocess.check_output(['nvidia-smi', '--query-gpu=memory.used,utilization.gpu', '--format=csv,noheader,nounits'], text=True)
    used, utilization = map(int, gpu.strip().split(','))
    command = "@(Get-Process -ErrorAction SilentlyContinue | Where-Object { $_.Path -like 'D:\\ITGlasses\\storage\\runtime\\*' } | Select-Object Id,CPU,WorkingSet64) | ConvertTo-Json -Compress"
    result = subprocess.run(['powershell', '-NoProfile', '-Command', command], capture_output=True, text=True)
    procs = json.loads(result.stdout) if result.stdout.strip() else []
    if isinstance(procs, dict): procs = [procs]
    return {'t': time.time(), 'phase': phase, 'free_ram_mib': round(memory.available / 2**20, 2), 'gpu_used_mib': used, 'gpu_utilization_percent': utilization, 'processes': procs}

def monitor():
    while not stop.wait(1):
        samples.append(measure())

def synthetic_files():
    from docx import Document
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    text = 'El procedimiento consiste en formatear el disco y reinstalar Windows.'
    doc = Document(); doc.add_heading('Synthetic Windows procedure', 1); doc.add_paragraph(text)
    docx = BytesIO(); doc.save(docx)
    writer = PdfWriter(); page = writer.add_blank_page(612, 792)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject(); stream.set_data(b'BT /F1 12 Tf 72 720 Td (Synthetic PDF evidence for procedure 43.) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    pdf = BytesIO(); writer.write(pdf)
    return [('windows.md', text.encode()), ('windows.txt', text.encode()), ('windows.docx', docx.getvalue()), ('procedure.pdf', pdf.getvalue())]

async def main():
    global phase
    with socket.socket() as sock:
        if sock.connect_ex(('127.0.0.1', 8001)) == 0:
            raise RuntimeError('Test port 8001 is occupied; stop this probe without changing the existing server.')
    baseline = measure()
    if baseline['free_ram_mib'] < 6144 or 6144 - baseline['gpu_used_mib'] < 3500:
        raise RuntimeError('Insufficient resource headroom for probe')
    assert baseline['processes'], 'Ollama process measurement unavailable'
    run = ROOT / 'storage/acceptance' / ('http-' + datetime.now().strftime('%Y%m%d-%H%M%S'))
    run.mkdir(parents=True)
    report = {'root': str(run), 'baseline': baseline, 'ingestion': []}
    samples.append(baseline)
    thread = threading.Thread(target=monitor, daemon=True); thread.start()
    server = None
    embedder = OllamaEmbeddings()
    try:
        phase = 'embedding_cpu_probe'
        started = time.perf_counter()
        await embedder.embed(['Synthetic CPU embedding resource measurement. ' * 30] * 16)
        report['embedding_seconds'] = round(time.perf_counter() - started, 3)
        async with httpx.AsyncClient(trust_env=False) as client:
            report['embedding_loaded'] = (await client.get('http://127.0.0.1:11434/api/ps')).json()
        samples.append(measure())
        assert all(m['size_vram'] == 0 for m in report['embedding_loaded']['models'])
        await embedder.unload()
        phase = 'http_ingestion'
        server = subprocess.Popen([sys.executable, __file__, '--serve', str(run)], cwd=ROOT, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        async with httpx.AsyncClient(base_url='http://127.0.0.1:8001', trust_env=False, timeout=120) as client:
            for _ in range(100):
                try:
                    if (await client.get('/health')).status_code == 200: break
                except httpx.ConnectError: pass
                await asyncio.sleep(0.1)
            else: raise RuntimeError('Backend startup failed')
            report['health'] = (await client.get('/health')).json()
            report['status'] = (await client.get('/api/v1/knowledge/status')).json()
            docs = []
            for i, (filename, raw) in enumerate(synthetic_files()):
                started = time.perf_counter()
                metadata = {'title': 'Formateo e instalación limpia de Windows' if i < 3 else 'Synthetic PDF procedure', 'scope': 'HTTP synthetic test', 'procedure_id': str(40+i), 'authorized': True}
                response = await client.post('/api/v1/documents', files={'file': (filename, raw)}, data={'metadata': json.dumps(metadata)})
                assert response.status_code == 201, response.text
                doc = response.json(); docs.append(doc)
                report['ingestion'].append({'format': Path(filename).suffix, 'bytes': len(raw), 'chunks': doc['chunk_count'], 'seconds': round(time.perf_counter()-started, 3)})
            phase = 'http_inference'
            started = time.perf_counter()
            response = await client.post('/api/v1/knowledge/chat', json={'message': '¿Qué describe el procedimiento 40 sobre formatear Windows?'})
            answer = response.json()
            assert answer['status'] == 'answered' and answer['sources'][0]['document_id'] == docs[0]['id'], answer
            report['query'] = answer
            report['query_seconds'] = round(time.perf_counter()-started, 3)
            response = await client.post('/api/v1/knowledge/chat', json={'message': 'procedimiento 43'})
            citation = response.json()['sources'][0]
            source = await client.get(citation['url'])
            assert source.status_code == 200 and source.json()['page'] == 1
            report['pdf_source_page'] = source.json()['page']
            withdrawn = await client.post('/api/v1/documents/' + docs[-1]['id'] + '/withdraw')
            assert withdrawn.status_code == 200
            assert (await client.get(citation['url'])).status_code == 404
            assert (await client.post('/api/v1/knowledge/chat', json={'message': ''})).status_code == 422
            assert (await client.post('/api/v1/documents', files={'file': ('invalid.exe', b'test')}, data={'metadata': json.dumps({'title':'Invalid'})})).status_code == 400
            report['error_checks'] = {'withdrawn_source': 404, 'empty_question': 422, 'unsupported_file': 400}
            report['listeners'] = subprocess.check_output(['netstat', '-ano', '-p', 'tcp'], text=True)
            report['listeners'] = [line.strip() for line in report['listeners'].splitlines() if ('127.0.0.1:8001 ' in line or '127.0.0.1:11434 ' in line) and 'LISTEN' in line]
            assert len(report['listeners']) == 2
            report['passed'] = True
    finally:
        if server is not None:
            server.terminate(); server.wait(timeout=10)
        stop.set(); thread.join(timeout=10)
        await embedder.unload()
        async with httpx.AsyncClient(trust_env=False, timeout=15) as client:
            await client.post('http://127.0.0.1:11434/api/generate', json={'model': 'qwen3:4b-instruct-2507-q4_K_M', 'keep_alive': 0})
        report['samples'] = samples
        report['minimum_free_ram_mib'] = min(s['free_ram_mib'] for s in samples)
        report['peak_total_gpu_mib'] = max(s['gpu_used_mib'] for s in samples)
        report['peak_ollama_working_set_mib'] = round(max(sum(p['WorkingSet64'] for p in s['processes']) for s in samples)/2**20, 2)
        probe = [s for s in samples if s['phase'] == 'embedding_cpu_probe']
        initial_cpu = {p['Id']: p['CPU'] for p in baseline['processes']}
        cpu_seconds = sum(max(p['CPU'] for s in probe for p in s['processes'] if p['Id'] == pid) - initial_cpu.get(pid, 0) for pid in {p['Id'] for s in probe for p in s['processes']})
        report['embedding_cpu_seconds_observed'] = round(cpu_seconds, 3)
        report['embedding_peak_working_set_mib'] = round(max(sum(p['WorkingSet64'] for p in s['processes']) for s in probe)/2**20, 2)
        (run / 'report.json').write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding='utf-8')
        print(json.dumps({k:v for k,v in report.items() if k not in {'samples','query'}}, ensure_ascii=False), flush=True)

if __name__ == '__main__':
    asyncio.run(main())
